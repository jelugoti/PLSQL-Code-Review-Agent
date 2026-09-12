from docx import Document


def read_standards_document(file_path):

    document = Document(file_path)

    content = []


    # ==================================================
    # Paragraphs
    # ==================================================

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            content.append({
                "type": "paragraph",
                "text": text
            })


    # ==================================================
    # Tables
    # ==================================================

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):

        rows = []

        for row in table.rows:

            cells = []

            for cell in row.cells:

                cells.append(
                    cell.text.strip()
                )

            rows.append(cells)


        content.append({
            "type": "table",
            "table_number": table_index,
            "rows": rows
        })


    return content